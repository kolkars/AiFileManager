from docx import Document
from openpyxl import Workbook

from ai_file_manager.extractors import default_registry


def test_initial_text_csv_docx_xlsx_extractors(tmp_path):
    registry = default_registry()
    txt = tmp_path / "a.txt"
    txt.write_text("plain", encoding="utf-8")
    csv = tmp_path / "a.csv"
    csv.write_text("name,value\nalpha,1", encoding="utf-8")
    docx = tmp_path / "a.docx"
    document = Document()
    document.add_paragraph("word text")
    document.save(docx)
    xlsx = tmp_path / "a.xlsx"
    workbook = Workbook()
    workbook.active.append(["sheet text", 7])
    workbook.save(xlsx)
    assert registry.extract(txt) == "plain"
    assert "alpha\t1" in registry.extract(csv)
    assert "word text" in registry.extract(docx)
    assert "sheet text\t7" in registry.extract(xlsx)

